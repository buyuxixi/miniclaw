"""Real extraction/storage with a counted offline model seam; no supplier requests."""
import csv
import hashlib
import io
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from miniclaw_web.attachment_uploads import AttachmentUploads
from miniclaw_web.document_analysis import DocumentAnalysis
from miniclaw_web.documents import DocumentService
from miniclaw_web.store import Store


class CountedModel:
    def __init__(self):
        self.requests = []
        self.hook = None

    def __call__(self, messages):
        value = json.loads(messages[-1]['content'])
        self.requests.append(value)
        if self.hook:
            override = self.hook(value)
            if override is not None:
                return override
        if isinstance(value['data'][0], dict) and 'id' in value['data'][0]:
            return json.dumps({'items': [{'id': item['id'], 'value': 'checked-' + item['id'][:12]}
                                         for item in value['data']]})
        return json.dumps({'summary': 'All supplied branches merged; unresolved differences retained.'})

    def mapped_ids(self):
        return [item['id'] for request in self.requests for item in request['data']
                if isinstance(item, dict) and 'id' in item]


class AnalysisTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.store = Store(Path(self.temp.name))
        self.documents = DocumentService(self.store, autostart=False)
        self.uploads = AttachmentUploads(self.store, self.documents)
        self.model = CountedModel()
        self.analysis = DocumentAnalysis(self.documents, call=self.model, autostart=False)

    def tearDown(self):
        self.analysis.shutdown()
        self.documents.shutdown()
        self.temp.cleanup()

    def document(self, rows=45, body='facts', name='facts.csv', raw=None):
        if raw is None:
            stream = io.StringIO(newline='')
            writer = csv.writer(stream)
            writer.writerow(['product', 'quantity', 'facts'])
            for number in range(rows):
                writer.writerow(['same-product' if number % 2 else f'product-{number}', number + 1, body])
            raw = stream.getvalue().encode()
        item = self.uploads.init('chat', name, len(raw), 'source-' + hashlib.sha256(raw).hexdigest()[:16])
        for number, offset in enumerate(range(0, len(raw), self.uploads.part_size)):
            part = raw[offset:offset + self.uploads.part_size]
            self.uploads.put_part(item['id'], 'chat', number, part, hashlib.sha256(part).hexdigest())
        self.uploads.complete(item['id'], 'chat')
        self.documents._execute(item['id'], 'chat')
        return item['id']

    def job(self, identifier, mode='scan', max_calls=128, request='analysis-one'):
        return self.analysis.create('chat', [identifier], mode, 'Read every complete unit; preserve references.', request, max_calls)

    def test_scan_covers_every_original_row_once_including_duplicate_business_values(self):
        identifier = self.document()
        expected = list(self.documents.iter_rows(identifier, 'chat'))
        job = self.job(identifier)
        self.analysis._run(job['id'], 'chat')
        status = self.analysis.status(job['id'], 'chat', limit=20)
        ids = self.model.mapped_ids()
        self.assertEqual(status['state'], 'complete')
        self.assertTrue(status['coverage']['complete'])
        self.assertEqual(status['coverage']['completed_units'], 45)
        self.assertEqual(set(ids), {row['id'] for row in expected})
        self.assertEqual(len(ids), len(set(ids)))
        results = [item for step in status['items'] for item in (step['output'] or {}).get('items', [])]
        self.assertEqual(len(results), 45)
        self.assertTrue(all('row_number' in item['source'] for item in results))
        self.assertEqual(status['sources'][0]['object_id'], self.store.get('attachments', identifier, 'chat')['object_id'])

    def test_long_table_row_is_mapped_as_one_complete_row(self):
        body = 'ROW_START-' + 'x' * 8000 + '-ROW_END'
        identifier = self.document(rows=1, body=body)
        self.assertGreater(self.documents.status(identifier, 'chat')['chunks'], 1)
        job = self.job(identifier)
        self.analysis._run(job['id'], 'chat')
        mapped = self.model.requests[0]['data']
        self.assertEqual(len(mapped), 1)
        self.assertEqual(mapped[0]['content']['facts'], body)
        self.assertTrue(self.analysis.status(job['id'], 'chat')['coverage']['complete'])

    def test_summarize_creates_multiple_reduce_levels_and_covers_every_chunk(self):
        identifier = self.document(rows=160, body='x' * 2000)
        expected = {chunk['id'] for chunk in self.documents.iter_chunks(identifier, 'chat')}
        job = self.job(identifier, mode='summarize')
        self.analysis._run(job['id'], 'chat')
        status = self.analysis.status(job['id'], 'chat')
        self.assertEqual(status['state'], 'complete')
        self.assertTrue(status['coverage']['complete'])
        self.assertEqual(set(self.model.mapped_ids()), expected)
        self.assertEqual(len(self.model.mapped_ids()), len(expected))
        with self.store.connect() as db:
            self.assertGreater(db.execute("SELECT count(*) FROM document_analysis_steps WHERE job_id=? AND phase='reduce'", (job['id'],)).fetchone()[0], 1)
        self.assertIn('summary', status['result'])

    def test_budget_pause_and_resume_never_replays_completed_units(self):
        identifier = self.document()
        job = self.job(identifier, max_calls=1)
        self.analysis._run(job['id'], 'chat')
        paused = self.analysis.status(job['id'], 'chat')
        self.assertEqual(paused['state'], 'paused')
        self.assertEqual(paused['calls'], 1)
        self.assertEqual(paused['coverage']['completed_units'], 20)
        self.assertFalse(paused['coverage']['complete'])
        previous_ids = list(self.model.mapped_ids())
        self.analysis.resume(job['id'], 'chat', 5)
        self.analysis._run(job['id'], 'chat')
        complete = self.analysis.status(job['id'], 'chat')
        self.assertTrue(complete['coverage']['complete'])
        self.assertGreater(complete['reserved_input_bytes'], paused['reserved_input_bytes'])
        self.assertEqual(self.model.mapped_ids()[:20], previous_ids)
        self.assertEqual(len(self.model.mapped_ids()), 45)
        self.assertEqual(len(set(self.model.mapped_ids())), 45)

    def test_input_budget_can_be_explicitly_raised_without_reset_or_completed_replay(self):
        identifier = self.document(rows=25, body='x' * 4000)
        job = self.analysis.create('chat', [identifier], 'scan', 'Read every complete row.',
                                   'input-budget-job', max_calls=128, max_input_bytes=64_000)
        self.analysis._run(job['id'], 'chat')
        paused = self.analysis.status(job['id'], 'chat')
        self.assertEqual(paused['state'], 'paused')
        self.assertLess(paused['calls'], 128)
        self.assertGreater(paused['reserved_input_bytes'], 0)
        previous_ids = list(self.model.mapped_ids())
        with self.assertRaises(ValueError):
            self.analysis.resume(job['id'], 'chat', 128, max_input_bytes=63_999)
        self.analysis.resume(job['id'], 'chat', 128, max_input_bytes=1024 * 1024)
        self.analysis._run(job['id'], 'chat')
        complete = self.analysis.status(job['id'], 'chat')
        self.assertTrue(complete['coverage']['complete'])
        self.assertGreater(complete['reserved_input_bytes'], paused['reserved_input_bytes'])
        self.assertEqual(self.model.mapped_ids()[:len(previous_ids)], previous_ids)
        self.assertEqual(len(self.model.mapped_ids()), 25)
        self.assertEqual(len(set(self.model.mapped_ids())), 25)

    def test_long_chinese_reduce_outputs_converge_within_request_budget(self):
        identifier = self.document(rows=100, body='x' * 4000)
        def long_output(payload):
            if isinstance(payload['data'][0], dict) and 'id' in payload['data'][0]:
                return json.dumps({'items': [{'id': item['id'], 'value': '中文结果' * 100}
                                             for item in payload['data']]}, ensure_ascii=False)
            return json.dumps({'summary': '归并结果' * 1500}, ensure_ascii=False)
        self.model.hook = long_output
        job = self.job(identifier, mode='summarize', max_calls=128)
        self.analysis._run(job['id'], 'chat')
        status = self.analysis.status(job['id'], 'chat')
        self.assertEqual(status['state'], 'complete')
        self.assertTrue(status['coverage']['complete'])
        self.assertLess(status['calls'], 128)
        self.assertEqual(len(status['result']['summary']), 6000)

    def test_invalid_or_incomplete_model_output_fails_without_automatic_replay(self):
        for number, output in enumerate(['true', '[]', '{"items":[]}']):
            identifier = self.document(rows=1, body='failure-' + str(number))
            self.model.hook = lambda payload, output=output: output
            job = self.job(identifier, request='invalid-output-' + str(number))
            before = len(self.model.requests)
            self.analysis._run(job['id'], 'chat')
            status = self.analysis.status(job['id'], 'chat')
            self.assertEqual(status['state'], 'failed')
            self.assertFalse(status['coverage']['complete'])
            self.assertEqual(len(self.model.requests), before + 1)
            self.assertEqual(status['steps'], {'failed': 1})

    def test_unknown_receipt_and_service_restart_require_explicit_retry(self):
        identifier = self.document(rows=1)
        def timeout(payload):
            raise TimeoutError('physical request accepted but receipt missing')
        self.model.hook = timeout
        job = self.job(identifier)
        self.analysis._run(job['id'], 'chat')
        self.assertEqual(self.analysis.status(job['id'], 'chat')['state'], 'unknown')
        self.analysis.shutdown()
        self.analysis = DocumentAnalysis(self.documents, call=self.model, autostart=False)
        self.analysis._pump()
        self.assertEqual(len(self.model.requests), 1)
        with self.assertRaises(ValueError):
            self.analysis.resume(job['id'], 'chat', 129)
        self.model.hook = None
        self.analysis.resume(job['id'], 'chat', 129, retry_unknown=True)
        self.analysis._run(job['id'], 'chat')
        complete = self.analysis.status(job['id'], 'chat')
        self.assertTrue(complete['coverage']['complete'])
        self.assertEqual(complete['calls'], 2)
        self.assertEqual(len(self.model.requests), 2)

    def test_cancel_during_request_keeps_late_known_result_and_resume_does_not_repay(self):
        identifier = self.document(rows=25)
        job = self.job(identifier)
        def cancel_once(payload):
            self.model.hook = None
            self.analysis.cancel(job['id'], 'chat')
        self.model.hook = cancel_once
        self.analysis._run(job['id'], 'chat')
        canceled = self.analysis.status(job['id'], 'chat')
        self.assertEqual(canceled['state'], 'cancelled')
        self.assertEqual(canceled['steps'], {'done': 1, 'queued': 1})
        self.analysis.resume(job['id'], 'chat', 128)
        self.analysis._run(job['id'], 'chat')
        self.assertTrue(self.analysis.status(job['id'], 'chat')['coverage']['complete'])
        self.assertEqual(len(self.model.mapped_ids()), 25)
        self.assertEqual(len(set(self.model.mapped_ids())), 25)

    def test_cross_owner_and_changed_request_are_rejected(self):
        identifier = self.document(rows=1)
        job = self.job(identifier)
        for function, args in ((self.analysis.status, (job['id'], 'other')),
                               (self.analysis.cancel, (job['id'], 'other')),
                               (self.analysis.resume, (job['id'], 'other', 128)),
                               (self.analysis.create, ('other', [identifier], 'scan', 'read', 'other-owner'))):
            with self.assertRaises(ValueError):
                function(*args)
        with self.assertRaises(ValueError):
            self.analysis.create('chat', [identifier], 'scan', 'changed question', 'analysis-one')
        self.assertEqual(self.job(identifier)['id'], job['id'])

    def test_reparse_changes_attempt_and_stops_original_manifest_before_model_call(self):
        identifier = self.document(raw=b'a,b\n1,2,3\n')
        self.assertEqual(self.documents.status(identifier, 'chat')['parse_status'], 'partial')
        job = self.job(identifier)
        source = job['sources'][0]
        self.documents.retry(identifier, 'chat')
        self.documents._execute(identifier, 'chat')
        self.assertGreater(self.documents.status(identifier, 'chat')['attempt'], source['attempt'])
        self.analysis._run(job['id'], 'chat')
        status = self.analysis.status(job['id'], 'chat')
        self.assertEqual(status['state'], 'failed')
        self.assertEqual(status['sources'], job['sources'])
        self.assertEqual(len(self.model.requests), 0)

    def test_partial_plan_is_not_committed_or_reused_after_iteration_failure(self):
        identifier = self.document()
        original = self.documents.iter_rows
        def interrupted(*args, **kwargs):
            for index, row in enumerate(original(*args, **kwargs)):
                if index == 21:
                    raise ValueError('transient source failure during planning')
                yield row
        job = self.job(identifier)
        with patch.object(self.documents, 'iter_rows', interrupted):
            self.analysis._run(job['id'], 'chat')
        self.assertEqual(self.analysis.status(job['id'], 'chat')['state'], 'failed')
        with self.store.connect() as db:
            self.assertEqual(db.execute('SELECT count(*) FROM document_analysis_steps WHERE job_id=?', (job['id'],)).fetchone()[0], 0)
        self.analysis.resume(job['id'], 'chat', 128)
        self.analysis._run(job['id'], 'chat')
        self.assertTrue(self.analysis.status(job['id'], 'chat')['coverage']['complete'])
        self.assertEqual(len(self.model.mapped_ids()), 45)

    def test_inspect_bounds_many_long_unicode_headers_and_preserves_column_pagination(self):
        from openpyxl import Workbook
        from miniclaw_web.document_tools import query
        workbook = Workbook(write_only=True)
        headers = [f'column-{column}-' + '中' * 380 for column in range(20)]
        for number in range(10):
            sheet = workbook.create_sheet(f'sheet-{number}')
            sheet.append(headers)
            sheet.append(['1'] * 20)
        buffer = io.BytesIO()
        workbook.save(buffer)
        identifier = self.document(name='wide.xlsx', raw=buffer.getvalue())
        result = query(self.documents, self.analysis, 'chat', {'action': 'inspect', 'attachment_id': identifier})
        self.assertEqual(result['parse_status'], 'ready')
        self.assertLessEqual(len(json.dumps(result, ensure_ascii=False).encode()), 48_000)
        self.assertTrue(result['metadata'].get('has_more_tables') or any(table.get('has_more_columns')
                        for table in result['metadata']['tables']))
        selected = query(self.documents, self.analysis, 'chat', {'action': 'inspect', 'attachment_id': identifier,
                         'table': 'sheet-9', 'offset': 19, 'limit': 1})
        self.assertEqual(selected['metadata']['tables'][0]['columns'], [headers[19]])
        names, table_offset = [], 0
        for _ in range(20):
            page = query(self.documents, self.analysis, 'chat', {'action': 'inspect', 'attachment_id': identifier,
                         'offset': table_offset, 'limit': 10})
            names.extend(table['name'] for table in page['metadata']['tables'])
            if not page['metadata']['has_more_tables']:
                break
            self.assertGreater(page['metadata']['next_table_offset'], table_offset)
            table_offset = page['metadata']['next_table_offset']
        self.assertEqual(names, [f'sheet-{number}' for number in range(10)])
        for name in names:
            columns, column_offset = [], 0
            for _ in range(20):
                page = query(self.documents, self.analysis, 'chat', {'action': 'inspect', 'attachment_id': identifier,
                             'table': name, 'offset': column_offset, 'limit': 20})
                schema = page['metadata']['tables'][0]
                columns.extend(schema['columns'])
                self.assertLessEqual(len(json.dumps(page, ensure_ascii=False).encode()), 48_000)
                if not schema['has_more_columns']:
                    break
                self.assertGreater(schema['next_column_offset'], column_offset)
                column_offset = schema['next_column_offset']
            self.assertEqual(columns, headers)

    def test_calculate_rejects_oversized_groups_without_truncating_full_statistics(self):
        from miniclaw_web.document_tools import query
        stream = io.StringIO(newline='')
        writer = csv.writer(stream)
        writer.writerow(['group', 'amount'])
        keys = [f'group-{number:02d}-' + '中文' * 2000 for number in range(20)]
        for number, key in enumerate(keys, 1):
            writer.writerow([key, number])
        identifier = self.document(name='large-groups.csv', raw=stream.getvalue().encode())
        original = self.documents.aggregate(identifier, 'chat', operation='sum', column='amount', group_by=['group'])
        self.assertTrue(original['complete'])
        self.assertEqual(original['examined_rows'], 20)
        self.assertEqual([item['group']['group'] for item in original['groups']], keys)
        args = {'action': 'calculate', 'attachment_id': identifier, 'operation': 'sum', 'column': 'amount'}
        with self.assertRaisesRegex(ValueError, 'group_by.*filters'):
            query(self.documents, self.analysis, 'chat', {**args, 'group_by': ['group']})
        total = query(self.documents, self.analysis, 'chat', args)
        self.assertTrue(total['complete'])
        self.assertEqual(total['examined_rows'], 20)
        self.assertEqual(total['groups'][0]['value'], '210')
        self.assertLessEqual(len(json.dumps(dict(success=True, **total), ensure_ascii=False).encode()), 48_000)
        filtered = query(self.documents, self.analysis, 'chat', {**args, 'group_by': ['group'],
                         'filters': [{'column': 'group', 'op': 'contains', 'value': 'group-00-'}]})
        self.assertTrue(filtered['complete'])
        self.assertEqual(filtered['examined_rows'], 20)
        self.assertEqual(filtered['matched_rows'], 1)
        self.assertEqual(filtered['groups'][0]['group']['group'], keys[0])
        self.assertEqual(filtered['groups'][0]['value'], '1')
        self.assertEqual(self.model.requests, [])

    def test_calculate_bounds_warning_display_and_preserves_row_locations_and_partial_coverage(self):
        from miniclaw_web.document_tools import query
        raw = b'group,amount\n' + b'group-a,1,extra\n' * 7
        identifier = self.document(name='warning-rows.csv', raw=raw)
        result = query(self.documents, self.analysis, 'chat',
                       {'action': 'calculate', 'attachment_id': identifier, 'operation': 'sum', 'column': 'amount'})
        self.assertFalse(result['complete'])
        self.assertEqual(result['examined_rows'], 7)
        self.assertEqual(result['groups'][0]['value'], '7')
        self.assertEqual(result['warning_count'], 7)
        self.assertTrue(result['warnings_truncated'])
        self.assertEqual(len(result['warnings']), 5)
        self.assertEqual([item['source']['line_start'] for item in result['warnings']], [2, 3, 4, 5, 6])
        self.assertTrue(all(item['source']['table'] == 'csv' for item in result['warnings']))
        self.assertEqual(self.model.requests, [])


if __name__ == '__main__':
    unittest.main()
