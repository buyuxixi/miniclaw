import io
import json
import threading
import time
import uuid
import zipfile
from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook
from pypdf import PdfReader, PdfWriter
from pypdf.generic import ArrayObject, DictionaryObject, NameObject, NumberObject, DecodedStreamObject

from miniclaw_web import documents
from miniclaw_web.documents import DocumentService, PARSER_VERSION
from miniclaw_web.store import Store


@pytest.fixture
def service(tmp_path):
    store = Store(tmp_path)
    result = DocumentService(store)
    yield result
    result.shutdown()


def add(service, name, raw, owner='chat', kind='document'):
    identifier = uuid.uuid4().hex
    obj = service.store.media.put_bytes(raw, 'application/octet-stream')
    service.store.put('attachments', identifier, dict(id=identifier, name=name, kind=kind,
                      bytes=len(raw), object_id=obj['id']), owner)
    service.enqueue(identifier, owner)
    return identifier


def parsed(service, identifier, owner='chat'):
    result = service.wait(identifier, owner, 10)
    assert result['parse_status'] not in ('queued', 'parsing'), result
    return result


def pdf_bytes(with_blank=True):
    writer = PdfWriter()
    page = writer.add_blank_page(width=200, height=200)
    font = DictionaryObject({NameObject('/Type'): NameObject('/Font'),
                             NameObject('/Subtype'): NameObject('/Type1'),
                             NameObject('/BaseFont'): NameObject('/Helvetica')})
    page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): writer._add_object(font)})})
    content = DecodedStreamObject()
    content.set_data(b'BT /F1 12 Tf 10 100 Td (document page one) Tj ET')
    page[NameObject('/Contents')] = writer._add_object(content)
    if with_blank:
        writer.add_blank_page(width=200, height=200)
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


def cmap_pdf_bytes(page_codes):
    """Real Identity-H CMap mappings from the reported CambriaMath failure."""
    writer = PdfWriter()
    cmap = DecodedStreamObject()
    cmap.set_data(b'''/CIDInit /ProcSet findresource begin
12 dict begin begincmap
/CIDSystemInfo << /Registry (Adobe) /Ordering (UCS) /Supplement 0 >> def
/CMapName /TestUnicode def /CMapType 2 def
1 begincodespacerange <0000> <FFFF> endcodespacerange
3 beginbfchar <0723> <D835> <0AD9> <DFD3> <0ADF> <DFCF> endbfchar
endcmap CMapName currentdict /CMap defineresource pop end end''')
    descendant = DictionaryObject({
        NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/CIDFontType2'),
        NameObject('/BaseFont'): NameObject('/CambriaMath'),
        NameObject('/CIDSystemInfo'): DictionaryObject({
            NameObject('/Registry'): __import__('pypdf').generic.TextStringObject('Adobe'),
            NameObject('/Ordering'): __import__('pypdf').generic.TextStringObject('Identity'),
            NameObject('/Supplement'): NumberObject(0)}),
    })
    font = DictionaryObject({
        NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type0'),
        NameObject('/BaseFont'): NameObject('/CambriaMath'), NameObject('/Encoding'): NameObject('/Identity-H'),
        NameObject('/DescendantFonts'): ArrayObject([writer._add_object(descendant)]),
        NameObject('/ToUnicode'): writer._add_object(cmap),
    })
    font_reference = writer._add_object(font)
    for codes in page_codes:
        page = writer.add_blank_page(width=200, height=200)
        page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): font_reference})})
        content = DecodedStreamObject()
        content.set_data(b'BT /F1 12 Tf 10 100 Td <' + codes.hex().encode('ascii') + b'> Tj ET')
        page[NameObject('/Contents')] = writer._add_object(content)
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


def workbook_bytes(formula=False):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Sales'
    sheet.append(['region', 'amount'])
    sheet.append(['east', .1])
    sheet.append(['east', '=1+1' if formula else .2])
    sheet.append(['west', 3])
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_text_large_document_preserves_every_piece_and_source(service):
    text = ('长段落甲' * 3000) + '\nsecond line\n'
    identifier = add(service, 'large.txt', text.encode())
    status = parsed(service, identifier)
    assert status['parse_status'] == 'ready'
    assert status['chunks'] > 1
    chunks = list(service.iter_chunks(identifier, 'chat'))
    assert ''.join(chunk['text'] for chunk in chunks) == text.replace('\n', '')
    assert all(len(chunk['text'].encode()) <= 16000 for chunk in chunks)
    assert chunks[0]['source']['line_start'] == 1
    assert chunks[-1]['source']['line_start'] == 2
    assert all(chunk['source']['object_id'] == status['object_id'] for chunk in chunks)
    assert all(chunk['source']['parser_version'] == PARSER_VERSION for chunk in chunks)
    assert 'text' not in service.store.get('attachments', identifier, 'chat')
    page = service.read(identifier, 'chat', limit=1)
    assert page['has_more'] and page['next_offset'] == 1
    second = service.read(identifier, 'chat', offset=page['next_offset'], limit=1)
    assert second['items'][0]['id'] != page['items'][0]['id']


def test_csv_decimal_full_scan_groups_filters_and_multiline_source(service):
    identifier = add(service, 'sales.csv', b'region,amount,note\neast,0.1,"two\nlines"\neast,0.2,x\nwest,2.00,y\n')
    status = parsed(service, identifier)
    assert status['parse_status'] == 'ready' and status['rows'] == 3
    result = service.aggregate(identifier, 'chat', 'sum', 'amount', group_by='region')
    assert {row['group']['region']: row['value'] for row in result['groups']} == {'east': '0.3', 'west': '2.00'}
    assert result['examined_rows'] == 3 and result['complete']
    average = service.aggregate(identifier, 'chat', 'avg', 'amount', filters=[dict(column='region', op='eq', value='east')])
    assert average['groups'][0]['value'] == '0.15'
    assert average['matched_rows'] == 2
    rows = list(service.iter_rows(identifier, 'chat'))
    assert rows[0]['source']['line_start'] == 2 and rows[0]['source']['line_end'] == 3
    assert rows[0]['cells']['note'] == 'two\nlines'
    assert service.search(identifier, 'chat', 'west')['retrieval_only']
    assert service.search(identifier, 'chat', 'west')['items'][0]['source']['row_number'] == 3


def test_numeric_invalid_cells_are_reported_and_never_claim_complete(service):
    identifier = add(service, 'values.csv', b'name,value\na,1\nb,not numeric\nc,\nd,2\n')
    assert parsed(service, identifier)['parse_status'] == 'ready'
    result = service.aggregate(identifier, 'chat', 'sum', 'value')
    assert result['groups'][0]['value'] == '3'
    assert result['groups'][0]['empty_rows'] == 1
    assert result['non_numeric_rows'] == 1 and not result['complete']
    assert result['non_numeric_examples'][0]['source']['line_start'] == 3
    assert service.aggregate(identifier, 'chat')['groups'][0]['value'] == '4'
    with pytest.raises(ValueError, match='非数值'):
        service.aggregate(identifier, 'chat', filters=[dict(column='value', op='gt', value=0)])


def test_word_body_order_headings_and_tables(service):
    word = Document()
    word.add_heading('Sales heading', 1)
    word.add_paragraph('Introduction')
    table = word.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = 'name'
    table.rows[0].cells[1].text = 'amount'
    cells = table.add_row().cells
    cells[0].text, cells[1].text = 'Alice', '12.5'
    word.add_paragraph('After table')
    output = io.BytesIO()
    word.save(output)
    identifier = add(service, 'report.docx', output.getvalue())
    status = parsed(service, identifier)
    assert status['parse_status'] == 'ready' and status['rows'] == 1
    chunks = list(service.iter_chunks(identifier, 'chat'))
    assert [chunk['text'] for chunk in chunks] == ['Sales heading', 'Introduction', 'name\tamount', 'name: Alice\namount: 12.5', 'After table']
    assert chunks[3]['source']['heading'] == 'Sales heading'
    assert chunks[3]['source']['row_number'] == 2
    assert service.aggregate(identifier, 'chat', 'sum', 'amount', table='table_1')['groups'][0]['value'] == '12.5'


def test_pdf_page_coverage_reports_scan_and_cannot_claim_ready(service):
    identifier = add(service, 'pages.pdf', pdf_bytes())
    status = parsed(service, identifier)
    assert status['parse_status'] == 'partial'
    assert status['coverage']['total_units'] == 2
    assert status['coverage']['units_done'] == 1 and not status['coverage']['complete']
    assert status['warnings'][0]['code'] == 'pdf_no_text'
    assert status['warnings'][0]['source']['page'] == 2
    chunks = list(service.iter_chunks(identifier, 'chat'))
    assert chunks[0]['source']['page'] == 1 and 'document page one' in chunks[0]['text']
    assert not chunks[0]['coverage_complete']


def test_text_pdf_without_gaps_is_ready(service):
    identifier = add(service, 'text.pdf', pdf_bytes(False))
    assert parsed(service, identifier)['coverage']['complete']


def test_real_pdf_cmap_surrogates_preserve_pairs_report_loss_and_continue_all_pages(service):
    pair = bytes.fromhex('07230ad9')
    # One valid pair, one lone low half, and one lone high half, separated by
    # normal text so impossible mappings must not be guessed or discarded.
    codes = 'math '.encode('utf-16-be') + pair + bytes.fromhex('0adf') + ' '.encode('utf-16-be') + bytes.fromhex('0723') + ' remainder'.encode('utf-16-be')
    raw = cmap_pdf_bytes(['before'.encode('utf-16-be'), codes, 'after target'.encode('utf-16-be')])
    extracted = PdfReader(io.BytesIO(raw)).pages[1].extract_text()
    assert '\ud835\udfd3' in extracted and '\udfcf' in extracted
    with pytest.raises(UnicodeEncodeError):
        extracted.encode('utf-8')
    identifier = add(service, 'math.pdf', raw)
    result = parsed(service, identifier)
    assert result['parse_status'] == 'partial' and result['parse_error'] is None
    assert result['coverage']['total_units'] == result['coverage']['units_seen'] == 3
    assert result['coverage']['units_done'] == 2 and not result['coverage']['stopped_early']
    assert result['metadata']['source_exhausted']
    assert result['metadata']['unicode_pairs_normalized'] == 1
    assert result['metadata']['unicode_replacement_chars'] == 2
    assert result['metadata']['unicode_affected_pages'] == 1
    assert result['warnings'][0]['code'] == 'pdf_unicode_replaced'
    assert result['warnings'][0]['source']['page'] == 2
    chunks = list(service.iter_chunks(identifier, 'chat'))
    assert len(chunks) == 3 and '𝟓' in chunks[1]['text'] and chunks[1]['text'].count('�') == 2
    assert 'remainder' in chunks[1]['text'] and 'after target' in chunks[2]['text']
    assert all(chunk['text'].encode('utf-8') for chunk in chunks)
    found = service.search(identifier, 'chat', 'target')
    assert found['items'][0]['source']['page'] == 3
    json.dumps(found, ensure_ascii=False).encode('utf-8')


def test_real_pdf_valid_surrogate_pair_is_lossless_and_ready(service):
    identifier = add(service, 'pair.pdf', cmap_pdf_bytes([bytes.fromhex('07230ad9')]))
    result = parsed(service, identifier)
    assert result['parse_status'] == 'ready' and result['coverage']['complete']
    assert result['warnings'] == [] and result['metadata']['unicode_pairs_normalized'] == 1
    assert service.read(identifier, 'chat')['items'][0]['text'] == '𝟓'


def test_pdf_normal_unicode_and_pair_at_chunk_boundary_remain_lossless():
    original = '正常😀𠮷𝟓' * 1000
    assert documents._normalize_pdf_unicode(original) == (original, 0, 0)
    normalized, replacements, pairs = documents._normalize_pdf_unicode('a' * 3999 + '\ud835\udfd3' + 'z')
    pieces = [text for text, _ in documents._pieces(normalized)]
    assert replacements == 0 and pairs == 1 and pieces == ['a' * 3999 + '𝟓', 'z']
    assert ''.join(pieces).encode('utf-8') == ('a' * 3999 + '𝟓z').encode('utf-8')


def test_retry_rebuilds_old_parse_with_current_version_and_new_source_ids(service, monkeypatch):
    monkeypatch.setattr(documents, 'PARSER_VERSION', 'local-v1')
    identifier = add(service, 'old.pdf', pdf_bytes(False))
    assert parsed(service, identifier)['parser_version'] == 'local-v1'
    old_ids = [chunk['id'] for chunk in service.iter_chunks(identifier, 'chat')]
    with service.store.connect() as db:
        db.execute("UPDATE document_parses SET state='partial' WHERE attachment_id=?", (identifier,))
    monkeypatch.setattr(documents, 'PARSER_VERSION', PARSER_VERSION)
    service.retry(identifier, 'chat')
    result = parsed(service, identifier)
    assert result['parser_version'] == PARSER_VERSION and result['attempt'] == 2
    chunks = list(service.iter_chunks(identifier, 'chat'))
    assert all(chunk['source']['parser_version'] == PARSER_VERSION for chunk in chunks)
    assert [chunk['id'] for chunk in chunks] != old_ids


def test_excel_values_and_missing_formula_cache(service):
    good = add(service, 'sales.xlsx', workbook_bytes())
    assert parsed(service, good)['parse_status'] == 'ready'
    assert service.aggregate(good, 'chat', 'sum', 'amount', table='Sales')['groups'][0]['value'] == '3.3'
    identifier = add(service, 'formula.xlsx', workbook_bytes(True))
    status = parsed(service, identifier)
    assert status['parse_status'] == 'partial' and status['rows'] == 3
    assert status['warnings'][0]['code'] == 'excel_formula_cache_missing'
    result = service.aggregate(identifier, 'chat', 'sum', 'amount', table='Sales')
    assert result['groups'][0]['value'] == '3.1' and not result['complete']
    assert result['groups'][0]['empty_rows'] == 1
    assert status['warnings'][0]['source']['sheet'] == 'Sales'
    assert status['warnings'][0]['source']['row_number'] == 3
    assert status['warnings'][0]['source']['formula_missing_count'] == 1
    assert not status['warnings'][0]['source']['formula_missing_columns_truncated']


def test_wide_excel_formula_warnings_have_bounded_locator_and_keep_all_cells(service):
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet('Wide formulas')
    columns = ['c' + str(index).zfill(5) for index in range(1, 16_385)]
    sheet.append(columns)
    sheet.append(['=1+1'] * len(columns))
    output = io.BytesIO()
    workbook.save(output)
    identifier = add(service, 'wide-formulas.xlsx', output.getvalue())
    status = parsed(service, identifier)
    assert status['parse_status'] == 'partial' and status['rows'] == 1
    assert status['coverage']['units_seen'] == 2 and status['coverage']['units_done'] == 1
    assert not status['coverage']['complete'] and not status['coverage']['stopped_early']
    warning = status['warnings'][0]
    assert warning['code'] == 'excel_formula_cache_missing'
    locator = warning['source']
    assert len(json.dumps(locator).encode()) < 2000
    assert locator['formula_missing_columns'] == list(range(1, 21))
    assert locator['formula_missing_count'] == len(columns)
    assert locator['formula_missing_columns_truncated']
    row = next(service.iter_rows(identifier, 'chat'))
    assert list(row['cells']) == columns and all(value is None for value in row['cells'].values())
    assert row['source']['formula_missing_count'] == len(columns)
    first_row_chunk = next(chunk for chunk in service.iter_chunks(identifier, 'chat')
                           if chunk['source']['row_number'] == 2)
    page = service.read(identifier, 'chat', offset=first_row_chunk['ordinal'], limit=1)
    assert page['items'][0]['id'] == first_row_chunk['id']
    assert page['returned_bytes'] <= 48_000 and page['has_more']
    found = service.search(identifier, 'chat', 'c00001: ', limit=1)
    assert found['items'][0]['id'] == first_row_chunk['id']
    assert found['returned_bytes'] <= 48_000


def test_excel_stale_dimensions_do_not_drop_rows(service):
    original = workbook_bytes()
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(original)) as source, zipfile.ZipFile(output, 'w') as target:
        for item in source.infolist():
            raw = source.read(item.filename)
            if item.filename == 'xl/worksheets/sheet1.xml':
                raw = raw.replace(b'<dimension ref="A1:B4"', b'<dimension ref="A1:A1"')
            target.writestr(item.filename, raw)
    identifier = add(service, 'dimensions.xlsx', output.getvalue())
    assert parsed(service, identifier)['rows'] == 3


def test_owner_is_enforced_on_every_data_interface(service):
    identifier = add(service, 'secret.csv', b'name,value\na,10\n')
    parsed(service, identifier)
    calls = [lambda: service.status(identifier, 'another'), lambda: service.read(identifier, 'another'),
             lambda: service.search(identifier, 'another', '10'), lambda: list(service.iter_chunks(identifier, 'another')),
             lambda: list(service.iter_rows(identifier, 'another')), lambda: service.aggregate(identifier, 'another'),
             lambda: service.cancel(identifier, 'another'), lambda: service.retry(identifier, 'another'),
             lambda: service.enqueue(identifier, 'another')]
    for call in calls:
        with pytest.raises(ValueError, match='不属于'):
            call()


def test_budget_limit_keeps_progress_but_marks_partial(service, monkeypatch):
    monkeypatch.setattr(documents, 'MAX_UNITS', 2)
    identifier = add(service, 'limit.txt', b'first\nsecond\nthird\nfourth\n')
    status = parsed(service, identifier)
    assert status['parse_status'] == 'partial' and status['chunks'] == 2
    assert status['coverage']['stopped_early'] and not status['coverage']['complete']
    assert status['parse_error']
    assert [chunk['text'] for chunk in service.iter_chunks(identifier, 'chat')] == ['first', 'second']


def test_corrupt_and_unsupported_documents_fail_explicitly(service):
    identifier = add(service, 'broken.pdf', b'not a PDF')
    result = parsed(service, identifier)
    assert result['parse_status'] == 'failed' and result['parse_error']
    with pytest.raises(ValueError, match='解析完成'):
        list(service.iter_chunks(identifier, 'chat'))
    with pytest.raises(ValueError, match='不支持'):
        add(service, 'legacy.xls', b'old file')


def test_cancel_retry_and_chunk_ids_stay_stable(tmp_path):
    service = DocumentService(Store(tmp_path), autostart=False)
    try:
        identifier = add(service, 'retries.txt', b'first\nsecond\n')
        assert service.cancel(identifier, 'chat')['parse_status'] == 'cancelled'
        service._execute(identifier, 'chat')
        assert service.status(identifier, 'chat')['chunks'] == 0
        service.retry(identifier, 'chat')
        service._execute(identifier, 'chat')
        first_ids = [chunk['id'] for chunk in service.iter_chunks(identifier, 'chat')]
        with service.store.connect() as db:
            db.execute("UPDATE document_parses SET state='partial' WHERE attachment_id=?", (identifier,))
        service.retry(identifier, 'chat')
        service._execute(identifier, 'chat')
        assert [chunk['id'] for chunk in service.iter_chunks(identifier, 'chat')] == first_ids
        assert service.status(identifier, 'chat')['attempt'] == 2
    finally:
        service.shutdown()


def test_restart_requeues_and_rebuilds_interrupted_parse(tmp_path):
    store = Store(tmp_path)
    original = DocumentService(store, autostart=False)
    identifier = add(original, 'resume.csv', b'name,value\na,1\nb,2\n')
    with store.connect() as db:
        db.execute("UPDATE document_parses SET state='parsing',attempt=1,chunk_count=1 WHERE attachment_id=?", (identifier,))
        db.execute('INSERT INTO document_chunks VALUES (?,?,?,?,?)', ('stale', identifier, 0, 'incomplete old value', '{}'))
    original.shutdown()
    resumed = DocumentService(Store(tmp_path))
    try:
        status = parsed(resumed, identifier)
        assert status['parse_status'] == 'ready' and status['attempt'] == 2 and status['rows'] == 2
        assert all(chunk['id'] != 'stale' for chunk in resumed.iter_chunks(identifier, 'chat'))
        assert resumed.aggregate(identifier, 'chat', 'sum', 'value')['groups'][0]['value'] == '3'
    finally:
        resumed.shutdown()


def test_late_worker_cannot_overwrite_cancelled_state(service, monkeypatch):
    entered, release = threading.Event(), threading.Event()
    def blocked(path, metadata):
        yield documents._unit('first', {'line_start': 1})
        entered.set()
        assert release.wait(5)
        yield documents._unit('second', {'line_start': 2})
    monkeypatch.setattr(documents, '_text_units', blocked)
    identifier = add(service, 'cancel.txt', b'a\nb')
    assert entered.wait(5)
    assert service.cancel(identifier, 'chat')['parse_status'] == 'cancelled'
    release.set()
    deadline = time.monotonic() + 5
    while identifier in service.futures and time.monotonic() < deadline:
        time.sleep(.01)
    status = service.status(identifier, 'chat')
    assert status['parse_status'] == 'cancelled' and status['chunks'] == 0


def test_only_two_parse_workers_and_durable_queue(service, monkeypatch):
    lock, release = threading.Lock(), threading.Event()
    counts = {'running': 0, 'peak': 0}
    def blocked(path, metadata):
        with lock:
            counts['running'] += 1
            counts['peak'] = max(counts['peak'], counts['running'])
        try:
            assert release.wait(5)
            yield documents._unit('parsed')
        finally:
            with lock:
                counts['running'] -= 1
    monkeypatch.setattr(documents, '_text_units', blocked)
    identifiers = [add(service, 'file' + str(index) + '.txt', b'source') for index in range(3)]
    deadline = time.monotonic() + 5
    while counts['running'] < 2 and time.monotonic() < deadline:
        time.sleep(.01)
    assert counts['running'] == 2
    with service.store.connect() as db:
        assert db.execute("SELECT count(*) FROM document_parses WHERE state='queued'").fetchone()[0] == 1
    release.set()
    assert service.wait_all(10)
    assert counts['peak'] == 2
    assert all(service.status(item, 'chat')['parse_status'] == 'ready' for item in identifiers)


def test_enqueue_idempotent_and_aggregate_rejects_arbitrary_operations(service):
    identifier = add(service, 'input.csv', b'name,value\na,1\n')
    status = parsed(service, identifier)
    same = service.enqueue(identifier, 'chat')
    assert same['attempt'] == status['attempt']
    with pytest.raises(ValueError, match='仅支持'):
        service.aggregate(identifier, 'chat', operation='DROP TABLE')
    with pytest.raises(ValueError, match='指定列'):
        service.aggregate(identifier, 'chat', operation='sum', column='missing')
    with pytest.raises(ValueError, match='筛选条件'):
        service.aggregate(identifier, 'chat', filters=[{'sql': '1=1'}])
    # SQL wildcard characters are searched literally.
    assert service.search(identifier, 'chat', '%')['items'] == []


def test_iterators_release_sqlite_before_consumers_write(service):
    identifier = add(service, 'write-consumer.csv', b'name,value\n' + b'a,1\n' * 130)
    assert parsed(service, identifier)['rows'] == 130
    for iterator in (service.iter_rows(identifier, 'chat'), service.iter_chunks(identifier, 'chat')):
        total = 0
        for item in iterator:
            with service.store.connect() as db:
                db.execute('INSERT OR REPLACE INTO bindings VALUES (?,?)', ('consumer-progress', json.dumps({'id': item['id']})))
            total += 1
        assert total >= 130


def test_iterators_detect_generation_change_between_pages(service):
    identifier = add(service, 'change.csv', b'name,value\n' + b'a,1\n' * 130)
    parsed(service, identifier)
    for method in (service.iter_rows, service.iter_chunks):
        iterator = method(identifier, 'chat')
        for _ in range(100):
            next(iterator)
        with service.store.connect() as db:
            db.execute('UPDATE document_parses SET attempt=attempt+1 WHERE attachment_id=?', (identifier,))
        with pytest.raises(ValueError, match='版本或状态已变化'):
            next(iterator)


def test_csv_duplicate_headers_and_extra_columns_never_drop_cells(service):
    identifier = add(service, 'headers.csv', b'a,a,extra_column_4\n1,2,3,4\n5,6,7\n')
    status = parsed(service, identifier)
    assert status['parse_status'] == 'partial' and status['rows'] == 2
    rows = list(service.iter_rows(identifier, 'chat'))
    assert rows[0]['cells'] == {'a': '1', 'a [2]': '2', 'extra_column_4': '3', 'extra_column_4 [2]': '4'}
    assert rows[1]['cells'] == {'a': '5', 'a [2]': '6', 'extra_column_4': '7'}
    assert len(status['warnings']) == 1


def test_large_word_heading_has_bounded_locator_without_losing_body(service):
    word = Document()
    heading = 'heading ' * 1500
    word.add_heading(heading, 1)
    word.add_paragraph('after heading')
    output = io.BytesIO()
    word.save(output)
    identifier = add(service, 'heading.docx', output.getvalue())
    assert parsed(service, identifier)['parse_status'] == 'ready'
    chunks = list(service.iter_chunks(identifier, 'chat'))
    assert ''.join(item['text'] for item in chunks[:-1]) == heading
    assert all(len(item['source']['heading']) == 500 for item in chunks)
    assert chunks[-1]['source']['heading_label_truncated']


def test_decimal_sum_is_exact_across_entire_accepted_precision_range(service):
    large = '9' * 1000 + '0'
    small = '0.' + '0' * 999 + '9' * 1000
    # A trailing zero has no significance and does not expand the decimal
    # coefficient bound, so express the large number with an exponent.
    raw = ('name,value\nlarge,' + '9' * 1000 + 'E1\nsmall,' + small + '\n').encode()
    identifier = add(service, 'precision.csv', raw)
    assert parsed(service, identifier)['parse_status'] == 'ready'
    result = service.aggregate(identifier, 'chat', 'sum', 'value')
    assert result['groups'][0]['value'] == large + '.' + small[2:]
    assert result['complete']


def test_csv_very_wide_record_is_bounded_before_allocating_all_cells(service):
    identifier = add(service, 'wide.csv', b'name,value\na,1\n' + b'a,' * (documents.MAX_UNIT_BYTES // 2 + 1) + b'\n')
    result = parsed(service, identifier)
    assert result['parse_status'] == 'partial'
    assert result['rows'] == 1 and result['coverage']['stopped_early']
    assert 'CSV记录超过1 MiB' in result['parse_error']


def test_search_multiple_terms_returns_stable_union_and_paginates_only_unique_hits(service):
    identifier = add(service, 'search.txt', b'before\nalpha one\nbetween\nbeta two\nalpha beta three\nafter\nbeta four\ntail\n')
    parsed(service, identifier)
    first = service.search(identifier, 'chat', queries=['alpha', 'beta', 'alpha'], limit=2)
    assert first['queries'] == ['alpha', 'beta']
    assert first['hit_ordinals'] == [1, 3]
    assert first['context_ordinals'] == [] and not first['context_truncated']
    assert first['hit_count'] == 2 and first['context_count'] == 0
    assert first['next_offset'] == 2 and first['has_more']
    second = service.search(identifier, 'chat', queries=['alpha', 'beta'], offset=first['next_offset'], limit=2)
    assert second['hit_ordinals'] == [4, 6]
    assert second['next_offset'] == 4 and not second['has_more']
    assert len({item['id'] for item in first['items'] + second['items']}) == 4
    legacy = service.search(identifier, 'chat', 'alpha')
    single = service.search(identifier, 'chat', queries=['alpha'])
    assert single['items'] == legacy['items'] and single['hit_ordinals'] == [1, 4]


def test_search_neighbors_are_ordered_deduplicated_and_do_not_advance_hit_pagination(service):
    identifier = add(service, 'context.txt', b'intro\nalpha\nbridge\nbeta\ntail\nalpha again\nend\n')
    parsed(service, identifier)
    first = service.search(identifier, 'chat', queries=['alpha', 'beta'], context=1, limit=2)
    assert first['hit_ordinals'] == [1, 3]
    assert first['context_ordinals'] == [0, 2, 4]
    assert [item['ordinal'] for item in first['items']] == [0, 1, 2, 3, 4]
    assert first['hit_count'] == 2 and first['context_count'] == 3
    assert first['next_offset'] == 2 and first['has_more'] and not first['context_truncated']
    assert first['retrieval_only'] and '不代表' in first['note']
    second = service.search(identifier, 'chat', queries=['alpha', 'beta'], context=1,
                            offset=first['next_offset'], limit=2)
    assert second['hit_ordinals'] == [5] and second['context_ordinals'] == [4, 6]
    assert second['next_offset'] == 3 and not second['has_more']
    assert second['items'][1]['source']['line_start'] == 6
    # A hit on a later page can also supply neighboring context without consuming it.
    adjacent = service.search(identifier, 'chat', queries=['alpha', 'bridge'], context=1, limit=1)
    assert adjacent['hit_ordinals'] == [1] and adjacent['context_ordinals'] == [0, 2]
    next_hit = service.search(identifier, 'chat', queries=['alpha', 'bridge'], offset=adjacent['next_offset'], limit=1)
    assert next_hit['hit_ordinals'] == [2]


def test_search_context_boundaries_and_no_hits_do_not_claim_truncation(service):
    identifier = add(service, 'edges.txt', b'start\ninside\nend\n')
    parsed(service, identifier)
    result = service.search(identifier, 'chat', 'start', context=2)
    assert result['hit_ordinals'] == [0] and result['context_ordinals'] == [1, 2]
    assert not result['context_truncated'] and not result['has_more']
    absent = service.search(identifier, 'chat', queries=['absent'], context=2, offset=3)
    assert absent['items'] == [] and absent['hit_ordinals'] == [] and absent['context_ordinals'] == []
    assert absent['hit_count'] == absent['context_count'] == 0
    assert absent['next_offset'] == 3 and not absent['has_more'] and not absent['context_truncated']


def test_search_budget_keeps_hits_before_context_and_marks_omitted_neighbors(service):
    long_text = '中' * 4000
    identifier = add(service, 'context-budget.txt',
                     ('\n'.join([long_text, 'target one', long_text, 'target two', long_text, 'target three']) + '\n').encode())
    parsed(service, identifier)
    result = service.search(identifier, 'chat', 'target', context=1, limit=2, max_bytes=20_000)
    assert result['hit_ordinals'] == [1, 3] and result['hit_count'] == 2
    assert result['context_ordinals'] == [0] and result['context_count'] == 1
    assert result['context_truncated'] and result['returned_bytes'] <= 20_000
    assert result['next_offset'] == 2 and result['has_more']
    assert [item['ordinal'] for item in result['items']] == [0, 1, 3]
    next_page = service.search(identifier, 'chat', 'target', context=1, offset=result['next_offset'], max_bytes=20_000)
    assert next_page['hit_ordinals'] == [5] and next_page['next_offset'] == 3
    assert next_page['context_ordinals'] == [4] and not next_page['context_truncated']


def test_search_hit_byte_budget_pages_without_skipping_original_matches(service):
    lines = ['target-' + str(number) + '-' + '中' * 3990 for number in range(4)]
    identifier = add(service, 'hit-budget.txt', ('\n'.join(lines) + '\n').encode())
    parsed(service, identifier)
    ordinals, offset = [], 0
    for _ in range(5):
        page = service.search(identifier, 'chat', queries=['target', '中'], context=2,
                              offset=offset, limit=4, max_bytes=20_000)
        assert page['hit_count'] == 1 and page['context_count'] == 0
        assert page['context_truncated'] and page['returned_bytes'] <= 20_000
        assert page['next_offset'] == offset + 1
        ordinals.extend(page['hit_ordinals'])
        offset = page['next_offset']
        if not page['has_more']:
            break
    assert ordinals == [0, 1, 2, 3]


def test_search_terms_escape_sql_wildcards_backslash_and_quotes_literally(service):
    identifier = add(service, 'literal.txt', b'plain text\nactual % sign\nunder_score\npath\\file\nalpha\n')
    parsed(service, identifier)
    union = service.search(identifier, 'chat', queries=['%', '_', '\\'])
    assert union['hit_ordinals'] == [1, 2, 3]
    assert service.search(identifier, 'chat', '%')['hit_ordinals'] == [1]
    assert service.search(identifier, 'chat', '_')['hit_ordinals'] == [2]
    assert service.search(identifier, 'chat', '\\')['hit_ordinals'] == [3]
    assert service.search(identifier, 'chat', queries=["' OR 1=1 --"])['items'] == []
    with pytest.raises(ValueError, match='不属于'):
        service.search(identifier, 'other', queries=['%', '_'], context=2)


@pytest.mark.parametrize('arguments', [
    {}, {'query': 'alpha', 'queries': ['beta']}, {'query': ''}, {'query': ' '},
    {'query': True}, {'query': 'x' * 501}, {'queries': []}, {'queries': 'alpha'},
    {'queries': [None]}, {'queries': ['alpha', '']}, {'queries': ['x'] * 6},
    {'queries': ['x' * 501]}, {'query': 'alpha', 'context': True},
    {'query': 'alpha', 'context': -1}, {'query': 'alpha', 'context': 3},
    {'query': 'alpha', 'context': 1.0}, {'query': 'alpha', 'context': '1'},
])
def test_search_rejects_malformed_terms_and_context_before_reading_data(service, arguments):
    with pytest.raises(ValueError):
        service.search('not-a-real-attachment', 'chat', **arguments)
