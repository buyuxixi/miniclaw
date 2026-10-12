"""Local document extraction and deterministic table operations.

SQLite owns parsing progress and bounded extracted chunks. The immutable media
object remains the source; this service never invokes a model or owns a chat.
"""
import csv
import hashlib
import json
import re
import threading
import time
import zipfile
from concurrent.futures import ThreadPoolExecutor
from decimal import Decimal, InvalidOperation, localcontext
from pathlib import Path

PARSER_VERSION = 'local-v2'
MAX_CHUNK_CHARS = 4000  # <=16 KiB UTF-8; conservative bytes, not an exact token count.
MAX_UNIT_BYTES = 1024 * 1024
MAX_EXTRACTED_BYTES = 200 * 1024 * 1024
MAX_UNITS = 500_000
MAX_CHUNKS = 100_000
MAX_ZIP_BYTES = 256 * 1024 * 1024
ACTIVE = ('queued', 'parsing')
TERMINAL = ('ready', 'partial', 'failed', 'cancelled')
EXTENSIONS = {'.pdf', '.docx', '.xlsx', '.csv', '.txt', '.md', '.markdown', '.srt'}


class ParseLimit(ValueError):
    pass


class ParseCancelled(Exception):
    pass


def _json(value):
    return json.dumps(value, ensure_ascii=False, separators=(',', ':'), allow_nan=False)


def _headers(values):
    """Blank and duplicate headers remain individually addressable."""
    result, used, occurrences = [], set(), {}
    for index, value in enumerate(values, 1):
        base = str(value).strip() if value is not None else ''
        base = base or 'column_' + str(index)
        candidate, occurrence = base, occurrences.get(base, 1)
        while candidate in used:
            occurrence += 1
            candidate = base + ' [' + str(occurrence) + ']'
        if len(candidate) > 500:
            raise ParseLimit('表头超过500字，未完成解析')
        used.add(candidate)
        occurrences[base] = occurrence
        result.append(candidate)
    return result


def _cell(value):
    if value is None:
        return None
    if hasattr(value, 'isoformat'):
        return value.isoformat()
    return str(value)


def _pieces(text):
    position, part = 0, 0
    while position < len(text):
        end = min(len(text), position + MAX_CHUNK_CHARS)
        if end < len(text):
            boundary = max(text.rfind('\n', end - 400, end), text.rfind(' ', end - 400, end))
            if boundary > position:
                end = boundary + 1
        part += 1
        yield text[position:end], {'part': part, 'char_start': position, 'char_end': end}
        position = end


def _unit(text='', source=None, cells=None, table=None, row_number=None, warnings=None, covered=True):
    return dict(text=text, source=source or {}, cells=cells, table=table,
                row_number=row_number, warnings=warnings or [], covered=covered)


def _check_zip(path):
    with zipfile.ZipFile(path) as package:
        files = package.infolist()
        if len(files) > 20_000 or sum(item.file_size for item in files) > MAX_ZIP_BYTES:
            raise ParseLimit('Office文件展开后超过256 MiB或文件项上限，未解析；原文件已保留')
        if any(item.flag_bits & 1 for item in files):
            raise ValueError('不支持加密的Office文件')


def _text_units(path, metadata):
    metadata.update(content_scope='text_lines', encoding='utf-8-sig')
    with path.open('r', encoding='utf-8-sig', newline='') as stream:
        line_number = 0
        while True:
            line = stream.readline(MAX_UNIT_BYTES + 1)
            if not line:
                break
            line_number += 1
            if len(line.encode('utf-8')) > MAX_UNIT_BYTES:
                raise ParseLimit('单行文本超过1 MiB，未完成解析')
            if '\x00' in line:
                raise ValueError('该文件包含二进制数据，不能作为文本解析')
            yield _unit(line.rstrip('\r\n'), {'line_start': line_number, 'line_end': line_number})


class _CSVLines:
    """Bound record bytes before csv.reader allocates a very wide row."""
    def __init__(self, stream):
        self.stream, self.record_bytes = stream, 0

    def __iter__(self):
        return self

    def __next__(self):
        line = self.stream.readline(MAX_UNIT_BYTES + 1)
        if not line:
            raise StopIteration
        self.record_bytes += len(line.encode('utf-8'))
        if self.record_bytes > MAX_UNIT_BYTES:
            raise ParseLimit('单个CSV记录超过1 MiB，未完成解析')
        return line


def _csv_units(path, metadata):
    metadata.update(content_scope='table_rows', encoding='utf-8-sig', empty_rows='excluded',
                    header_policy='first_nonempty_row')
    # csv.field_size_limit is process global; use one fixed bounded policy.
    csv.field_size_limit(MAX_UNIT_BYTES)
    with path.open('r', encoding='utf-8-sig', newline='') as stream:
        sample = stream.read(8192)
        stream.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=',;\t|')
        except csv.Error:
            dialect = csv.excel
        lines = _CSVLines(stream)
        reader = csv.reader(lines, dialect)
        previous_line, headers, data_number = 0, None, 0
        for values in reader:
            lines.record_bytes = 0
            start_line, previous_line = previous_line + 1, reader.line_num
            source = dict(table='csv', line_start=start_line, line_end=reader.line_num)
            if any('\x00' in value for value in values):
                raise ValueError('CSV含二进制数据，不能作为表格解析')
            if not values or not any(value.strip() for value in values):
                yield _unit(source=source)
                continue
            if headers is None:
                headers = _headers(values)
                metadata['tables'] = [{'name': 'csv', 'columns': list(headers), 'header_row': reader.line_num}]
                yield _unit('\t'.join(headers), {**source, 'header': True})
                continue
            if len(values) != len(headers):
                # Preserve all cells, while explicitly reporting the changed schema.
                width = max(len(values), len(headers))
                row_headers = _headers(headers + ['extra_column_' + str(i + 1) for i in range(len(headers), width)])
                known = metadata['tables'][0]['columns']
                for column in row_headers:
                    if column not in known:
                        known.append(column)
                warning = [{'code': 'csv_column_mismatch', 'source': source,
                            'message': '该行列数与表头不同，保留额外列及缺失空值'}]
            else:
                row_headers, warning = headers, []
            cells = {key: values[i] if i < len(values) else None for i, key in enumerate(row_headers)}
            data_number += 1
            yield _unit('\n'.join(key + ': ' + (value or '') for key, value in cells.items()),
                        {**source, 'row_number': data_number}, cells, 'csv', data_number, warning,
                        covered=not warning)
        if headers is None:
            metadata['tables'] = []


def _pdf_units(path, metadata):
    from pypdf import PdfReader
    metadata.update(content_scope='pdf_page_text', ocr_supported=False)
    with path.open('rb') as stream:
        reader = PdfReader(stream)
        if reader.is_encrypted:
            raise ValueError('不支持加密PDF，请先导出未加密版本')
        metadata['total_units'] = len(reader.pages)
        for page_number, page in enumerate(reader.pages, 1):
            source = {'page': page_number}
            try:
                text = page.extract_text() or ''
            except Exception:
                yield _unit(source=source, covered=False, warnings=[dict(
                    code='pdf_page_failed', source=source, message='该页文字提取失败')])
                continue
            text, replacements, pairs = _normalize_pdf_unicode(text)
            if pairs:
                metadata['unicode_pairs_normalized'] = metadata.get('unicode_pairs_normalized', 0) + pairs
            warning = []
            if replacements:
                metadata['unicode_replacement_chars'] = metadata.get('unicode_replacement_chars', 0) + replacements
                metadata['unicode_affected_pages'] = metadata.get('unicode_affected_pages', 0) + 1
                warning.append(dict(code='pdf_unicode_replaced', source=source,
                                    message=f'该页有{replacements}个无法识别的编码字符，已用�占位，其余文字继续读取；数学符号等内容需核对原件'))
            if not text.strip():
                yield _unit(source=source, covered=False, warnings=warning + [dict(
                    code='pdf_no_text', source=source, message='该页无可提取文字，可能为扫描页或空白页；未执行OCR')])
            else:
                yield _unit(text, source, warnings=warning, covered=not replacements)


_PDF_SURROGATES = re.compile(r'[\ud800-\udbff][\udc00-\udfff]|[\ud800-\udfff]')


def _normalize_pdf_unicode(text):
    """CMaps can map separate glyphs to UTF-16 halves; keep valid pairs losslessly."""
    replacements, pairs = 0, 0

    def replace(match):
        nonlocal replacements, pairs
        value = match.group()
        if len(value) == 2:
            pairs += 1
            return chr(0x10000 + ((ord(value[0]) - 0xD800) << 10) + ord(value[1]) - 0xDC00)
        replacements += 1
        return '\ufffd'

    return _PDF_SURROGATES.sub(replace, text), replacements, pairs


def _docx_units(path, metadata):
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    _check_zip(path)
    document = Document(path)
    metadata.update(content_scope='word_body_paragraphs_and_tables',
                    scope_notes='仅正文段落和表格；不含页眉页脚、批注、脚注、文本框和图像OCR',
                    header_policy='first_row_of_each_table', tables=[])
    heading, heading_truncated, paragraph_number, table_number = '', False, 0, 0
    for element in document.element.body.iterchildren():
        if element.tag.endswith('}p'):
            paragraph_number += 1
            paragraph = Paragraph(element, document)
            text = paragraph.text
            if paragraph.style is not None and paragraph.style.name.startswith('Heading'):
                heading, heading_truncated = text[:500], len(text) > 500
            source = {'paragraph': paragraph_number, 'heading': heading, 'heading_label_truncated': heading_truncated}
            drawings = bool(element.xpath('.//w:drawing | .//w:pict'))
            warning = [dict(code='word_image_not_read', source=source,
                            message='正文含图像，未提取图像文字')] if drawings else []
            yield _unit(text, source, warnings=warning, covered=not drawings)
        elif element.tag.endswith('}tbl'):
            table_number += 1
            name = 'table_' + str(table_number)
            table = Table(element, document)
            headers = None
            for row_number, row in enumerate(table.rows, 1):
                values = [cell.text for cell in row.cells]
                source = dict(table=name, row_number=row_number, heading=heading, heading_label_truncated=heading_truncated)
                if headers is None:
                    headers = _headers(values)
                    metadata['tables'].append(dict(name=name, columns=headers, header_row=1))
                    yield _unit('\t'.join(headers), {**source, 'header': True})
                    continue
                row_headers = _headers(headers + ['extra_column_' + str(i + 1) for i in range(len(headers), len(values))])
                cells = {key: values[i] if i < len(values) else None for i, key in enumerate(row_headers)}
                warning = []
                if len(values) != len(headers):
                    warning = [dict(code='word_table_columns', source=source,
                                    message='表格行与首行列数不同，保留原始单元格')]
                if row._tr.xpath('.//w:tbl | .//w:drawing | .//w:pict | .//w:gridSpan | .//w:vMerge'):
                    warning.append(dict(code='word_table_complex_content', source=source,
                                        message='该行含合并单元格、嵌套表格或图像，展开结果需核对'))
                yield _unit('\n'.join(key + ': ' + (value or '') for key, value in cells.items()),
                            source, cells, name, row_number, warning, covered=not warning)


def _xlsx_units(path, metadata):
    from openpyxl import load_workbook
    _check_zip(path)
    # Content-addressed object paths intentionally have no extension. Pass streams
    # so openpyxl validates the package rather than rejecting its local filename.
    cached_stream = path.open('rb')
    formula_stream = path.open('rb')
    cached = formulas = None
    try:
        cached = load_workbook(cached_stream, read_only=True, data_only=True, keep_links=False)
        formulas = load_workbook(formula_stream, read_only=True, data_only=False, keep_links=False)
        metadata.update(content_scope='workbook_cell_values', tables=[], empty_rows='excluded',
                        header_policy='first_nonempty_row_of_each_sheet',
                        scope_notes='包含隐藏行和工作表；读取已有公式缓存，不运行Excel公式，不含图表和图像OCR')
        for sheet in cached.worksheets:
            formula_sheet, headers = formulas[sheet.title], None
            # Workbook dimension hints can be stale, so consume the XML stream.
            sheet.reset_dimensions()
            formula_sheet.reset_dimensions()
            for row_number, (row, formula_row) in enumerate(zip(sheet.iter_rows(), formula_sheet.iter_rows()), 1):
                values = [_cell(cell.value) for cell in row]
                source = dict(table=sheet.title, sheet=sheet.title, row_number=row_number)
                missing = [index for index, cell in enumerate(formula_row)
                           if cell.data_type == 'f' and (index >= len(row) or row[index].value is None)]
                errors = [index for index, cell in enumerate(row) if cell.data_type == 'e']
                if not any(value not in (None, '') for value in values) and not missing:
                    yield _unit(source=source)
                    continue
                width = max(len(values), len(formula_row))
                values += [None] * (width - len(values))
                warning = []
                if missing:
                    # Locators accompany every chunk and analysis request. Keep
                    # a sample plus the exact count; retain all original cells.
                    source['formula_missing_columns'] = [index + 1 for index in missing[:20]]
                    source['formula_missing_count'] = len(missing)
                    source['formula_missing_columns_truncated'] = len(missing) > 20
                    warning.append(dict(code='excel_formula_cache_missing', source=source,
                                        message='公式无缓存值，未计算；统计不能作为全量数值结果'))
                if errors:
                    warning.append(dict(code='excel_cell_error', source=source, message='工作表含Excel错误值'))
                if headers is None:
                    headers = _headers(values)
                    metadata['tables'].append(dict(name=sheet.title, columns=headers, header_row=row_number))
                    yield _unit('\t'.join(headers), {**source, 'header': True}, warnings=warning,
                                covered=not warning)
                    continue
                if len(values) > len(headers):
                    additions = ['column_' + str(index + 1) for index in range(len(headers), len(values))]
                    headers = _headers(headers + additions)
                    metadata['tables'][-1]['columns'] = headers
                cells = {key: values[i] if i < len(values) else None for i, key in enumerate(headers)}
                yield _unit('\n'.join(key + ': ' + (value or '') for key, value in cells.items()),
                            source, cells, sheet.title, row_number, warning, covered=not warning)
    finally:
        if cached is not None:
            cached.close()
        if formulas is not None:
            formulas.close()
        cached_stream.close()
        formula_stream.close()


class DocumentService:
    def __init__(self, store, max_workers=2, autostart=True):
        self.store = store
        self.max_workers = max(1, min(2, int(max_workers)))
        self.executor = ThreadPoolExecutor(max_workers=self.max_workers, thread_name_prefix='miniclaw-documents')
        self.lock = threading.RLock()
        self.futures = {}
        self.closed = False
        self.autostart = autostart
        # Local extraction has no external effects and is safe to restart.
        with store.connect() as db:
            db.execute("UPDATE document_parses SET state='queued',error=NULL,updated_at=? WHERE state='parsing'", (time.time(),))
        if autostart:
            self._pump()

    @staticmethod
    def _validate_owner(owner):
        if not isinstance(owner, str) or not re.fullmatch(r'[\w-]{1,128}', owner):
            raise ValueError('会话标识无效')

    def _record(self, attachment_id, owner, db=None):
        self._validate_owner(owner)
        if not isinstance(attachment_id, str) or not re.fullmatch(r'[\w-]{1,128}', attachment_id):
            raise ValueError('附件标识无效')
        if db is None:
            with self.store.connect() as connection:
                return self._record(attachment_id, owner, connection)
        keys = ('attachment_id', 'owner', 'object_id', 'parser_version', 'state', 'attempt',
                'chunk_count', 'row_count', 'units_seen', 'units_done', 'warnings', 'metadata',
                'error', 'created_at', 'updated_at')
        row = db.execute('SELECT ' + ','.join(keys) + ' FROM document_parses WHERE attachment_id=? AND owner=?',
                         (attachment_id, owner)).fetchone()
        if row is None:
            raise ValueError('文档不存在或不属于当前对话')
        record = dict(zip(keys, row))
        record['warnings'], record['metadata'] = json.loads(record['warnings']), json.loads(record['metadata'])
        return record

    @staticmethod
    def _public(record):
        metadata = record['metadata']
        state = record['state']
        coverage = dict(units_seen=record['units_seen'], units_done=record['units_done'],
                        total_units=metadata.get('total_units'), complete=state == 'ready',
                        stopped_early=bool(metadata.get('stopped_early')), warning_count=metadata.get('warning_count', 0))
        return dict(attachment_id=record['attachment_id'], object_id=record['object_id'],
                    parser_version=record['parser_version'], parse_status=state, parse_error=record['error'],
                    chunks=record['chunk_count'], rows=record['row_count'], warnings=record['warnings'],
                    coverage=coverage, metadata=metadata, attempt=record['attempt'], updated_at=record['updated_at'])

    def status(self, attachment_id, owner):
        return self._public(self._record(attachment_id, owner))

    def enqueue(self, attachment_id, owner):
        self._validate_owner(owner)
        record = self.store.get('attachments', attachment_id, owner)
        if record.get('kind') == 'image':
            return dict(attachment_id=attachment_id, parse_status='ready', parse_error=None,
                        chunks=0, rows=0, warnings=[], coverage={'complete': True, 'content_scope': 'image'})
        if Path(record.get('name', '')).suffix.lower() not in EXTENSIONS:
            raise ValueError('不支持该文档类型，仅支持PDF、DOCX、XLSX、CSV及UTF-8文本')
        if not record.get('object_id'):
            raise ValueError('附件缺少原始文件对象')
        with self.lock:
            with self.store.connect() as db:
                existing = db.execute('SELECT object_id,owner FROM document_parses WHERE attachment_id=?',
                                      (attachment_id,)).fetchone()
                if existing and existing != (record['object_id'], owner):
                    raise ValueError('附件原始版本发生变化，请上传为新附件')
                if not existing:
                    pending = db.execute("SELECT count(*) FROM document_parses WHERE state IN ('queued','parsing')").fetchone()[0]
                    if pending >= 100:
                        raise ValueError('文档解析队列已满，请等待已有任务完成')
                    now = time.time()
                    db.execute('INSERT INTO document_parses '
                               '(attachment_id,owner,object_id,parser_version,state,created_at,updated_at) VALUES (?,?,?,?,?,?,?)',
                               (attachment_id, owner, record['object_id'], PARSER_VERSION, 'queued', now, now))
            if self.autostart:
                self._pump()
        return self.status(attachment_id, owner)

    def _pump(self):
        with self.lock:
            if self.closed:
                return
            with self.store.connect() as db:
                candidates = db.execute("SELECT attachment_id,owner FROM document_parses WHERE state='queued' ORDER BY created_at,attachment_id LIMIT ?",
                                        (self.max_workers + len(self.futures),)).fetchall()
            for identifier, owner in candidates:
                if len(self.futures) >= self.max_workers:
                    break
                if identifier in self.futures:
                    continue
                future = self.executor.submit(self._execute, identifier, owner)
                self.futures[identifier] = future
                future.add_done_callback(lambda completed, item=identifier: self._finished(item))

    def _finished(self, identifier):
        with self.lock:
            self.futures.pop(identifier, None)
            if self.autostart:
                self._pump()

    def _commit(self, record, batch, counters, warnings, metadata, final=None, error=None):
        with self.store.connect() as db:
            current = db.execute('SELECT state,attempt FROM document_parses WHERE attachment_id=?',
                                 (record['attachment_id'],)).fetchone()
            if current != ('parsing', record['attempt']):
                raise ParseCancelled()
            for chunk, row in batch:
                if row is not None:
                    db.execute('INSERT INTO document_rows VALUES (?,?,?,?,?)', row)
                for values in chunk:
                    db.execute('INSERT INTO document_chunks VALUES (?,?,?,?,?)', values)
            db.execute('UPDATE document_parses SET state=?,chunk_count=?,row_count=?,units_seen=?,units_done=?,warnings=?,metadata=?,error=?,updated_at=? '
                       'WHERE attachment_id=? AND attempt=?',
                       (final or 'parsing', counters['chunks'], counters['rows'], counters['seen'], counters['done'],
                        _json(warnings), _json(metadata), error, time.time(), record['attachment_id'], record['attempt']))
        batch.clear()

    def _execute(self, attachment_id, owner):
        record = None
        batch, warnings, metadata = [], [], {}
        counters = dict(chunks=0, rows=0, seen=0, done=0, bytes=0, warning_count=0)
        try:
            with self.lock:
                with self.store.connect() as db:
                    record = self._record(attachment_id, owner, db)
                    if record['state'] != 'queued':
                        return
                    record['attempt'] += 1
                    record['parser_version'] = PARSER_VERSION
                    db.execute("UPDATE document_parses SET state='parsing',attempt=?,parser_version=?,chunk_count=0,row_count=0,units_seen=0,units_done=0,warnings='[]',metadata='{}',error=NULL,updated_at=? WHERE attachment_id=? AND state='queued'",
                               (record['attempt'], record['parser_version'], time.time(), attachment_id))
                    db.execute('DELETE FROM document_chunks WHERE attachment_id=?', (attachment_id,))
                    db.execute('DELETE FROM document_rows WHERE attachment_id=?', (attachment_id,))
            attachment = self.store.get('attachments', attachment_id, owner)
            path = self.store.media.materialize(record['object_id'])
            extension = Path(attachment['name']).suffix.lower()
            parser = {'.pdf': _pdf_units, '.docx': _docx_units, '.xlsx': _xlsx_units, '.csv': _csv_units}.get(extension, _text_units)
            metadata.update(name=attachment['name'], extracted_bytes=0, warning_count=0,
                            chunk_max_utf8_bytes=MAX_CHUNK_CHARS * 4, token_estimation='utf8_bytes_conservative_upper_bound')
            batch_bytes = 0
            for unit in parser(path, metadata):
                if not batch:
                    with self.store.connect() as db:
                        current = db.execute('SELECT state,attempt FROM document_parses WHERE attachment_id=?', (attachment_id,)).fetchone()
                    if current != ('parsing', record['attempt']) or self.closed:
                        raise ParseCancelled()
                text = unit['text']
                size = len(text.encode('utf-8'))
                row_json = _json(unit['cells']) if unit['cells'] is not None else None
                if size > MAX_UNIT_BYTES or (row_json and len(row_json.encode('utf-8')) > MAX_UNIT_BYTES):
                    raise ParseLimit('单个段落、页或表格行超过1 MiB，未完成解析')
                if counters['seen'] >= MAX_UNITS or counters['bytes'] + size > MAX_EXTRACTED_BYTES:
                    raise ParseLimit('达到50万解析单位或200 MiB提取预算，未完成解析')
                parts = list(_pieces(text)) if text.strip() else []
                if counters['chunks'] + len(parts) > MAX_CHUNKS:
                    raise ParseLimit('达到10万个文字块上限，未完成解析')
                counters['seen'] += 1
                counters['done'] += int(unit['covered'])
                counters['bytes'] += size
                counters['warning_count'] += len(unit['warnings'])
                warnings.extend(unit['warnings'][:max(0, 100 - len(warnings))])
                source = dict(object_id=record['object_id'], parser_version=record['parser_version'],
                              attachment_id=attachment_id, **unit['source'])
                chunks = []
                for part, locator in parts:
                    ordinal = counters['chunks']
                    chunk_id = hashlib.sha256((attachment_id + ':' + record['object_id'] + ':' + record['parser_version'] + ':' + str(ordinal)).encode()).hexdigest()
                    chunks.append((chunk_id, attachment_id, ordinal, part, _json({**source, **locator})))
                    counters['chunks'] += 1
                row = None
                if row_json is not None:
                    row = (attachment_id, unit['table'], unit['row_number'], row_json, _json(source))
                    counters['rows'] += 1
                batch.append((chunks, row))
                batch_bytes += size + (len(row_json.encode('utf-8')) if row_json else 0)
                metadata.update(extracted_bytes=counters['bytes'], warning_count=counters['warning_count'])
                if len(batch) >= 64 or batch_bytes >= 512 * 1024:
                    self._commit(record, batch, counters, warnings, metadata)
                    batch_bytes = 0
            if not counters['seen']:
                warnings.append(dict(code='document_empty', message='文件没有可解析内容'))
                counters['warning_count'] += 1
            metadata.update(warning_count=counters['warning_count'], source_exhausted=True)
            metadata.setdefault('total_units', counters['seen'])
            final = 'partial' if counters['warning_count'] or counters['done'] != counters['seen'] else 'ready'
            self._commit(record, batch, counters, warnings, metadata, final)
        except ParseCancelled:
            # Cancellation/restart owns the state; a late worker must not overwrite it.
            return
        except Exception as error:
            if record is None:
                return
            message = str(error) if isinstance(error, (ValueError, UnicodeError, csv.Error)) else '文件解析失败，原文件已保留（' + type(error).__name__ + '）'
            metadata.update(stopped_early=True, warning_count=counters['warning_count'], extracted_bytes=counters['bytes'])
            try:
                self._commit(record, batch, counters, warnings, metadata,
                             'partial' if counters['chunks'] or counters['rows'] else 'failed', message[:1000])
            except ParseCancelled:
                pass

    def retry(self, attachment_id, owner):
        with self.lock:
            record = self._record(attachment_id, owner)
            if record['state'] not in ('failed', 'partial', 'cancelled'):
                raise ValueError('仅失败、部分完成或取消的文档可重试')
            if attachment_id in self.futures:
                raise ValueError('原解析仍在退出，请稍后重试')
            with self.store.connect() as db:
                db.execute("UPDATE document_parses SET state='queued',error=NULL,updated_at=? WHERE attachment_id=? AND owner=?",
                           (time.time(), attachment_id, owner))
            if self.autostart:
                self._pump()
        return self.status(attachment_id, owner)

    def cancel(self, attachment_id, owner):
        with self.lock:
            self._record(attachment_id, owner)
            with self.store.connect() as db:
                db.execute("UPDATE document_parses SET state='cancelled',error='用户已取消解析',updated_at=? WHERE attachment_id=? AND owner=? AND state IN ('queued','parsing')",
                           (time.time(), attachment_id, owner))
        return self.status(attachment_id, owner)

    def wait(self, attachment_id, owner, seconds=15):
        deadline = time.monotonic() + seconds
        while time.monotonic() < deadline:
            result = self.status(attachment_id, owner)
            if result['parse_status'] not in ACTIVE:
                return result
            time.sleep(.02)
        return self.status(attachment_id, owner)

    def wait_all(self, seconds=30):
        deadline = time.monotonic() + seconds
        while time.monotonic() < deadline:
            with self.store.connect() as db:
                active = db.execute("SELECT count(*) FROM document_parses WHERE state IN ('queued','parsing')").fetchone()[0]
            if not active:
                return True
            time.sleep(.02)
        return False

    def shutdown(self, wait=True):
        with self.lock:
            self.closed = True
        self.executor.shutdown(wait=wait)

    @staticmethod
    def _page(offset, limit, max_bytes):
        if type(offset) is not int or offset < 0 or type(limit) is not int or not 1 <= limit <= 50:
            raise ValueError('分页参数无效，limit需为1至50')
        if type(max_bytes) is not int or not 20_000 <= max_bytes <= 100_000:
            raise ValueError('每页文字预算需为20000至100000字节')

    @staticmethod
    def _chunk(row):
        return dict(id=row[0], ordinal=row[1], text=row[2], source=json.loads(row[3]),
                    conservative_token_budget=len(row[2].encode('utf-8')))

    def read(self, attachment_id, owner, offset=0, limit=10, max_bytes=48_000):
        self._page(offset, limit, max_bytes)
        status = self.status(attachment_id, owner)
        chunks, used = [], 0
        with self.store.connect() as db:
            rows = db.execute('SELECT id,ordinal,text,source FROM document_chunks WHERE attachment_id=? AND ordinal>=? ORDER BY ordinal LIMIT ?',
                              (attachment_id, offset, limit)).fetchall()
        for row in rows:
            chunk = self._chunk(row)
            cost = len(_json(chunk).encode('utf-8'))
            if chunks and used + cost > max_bytes:
                break
            chunks.append(chunk)
            used += cost
        next_offset = chunks[-1]['ordinal'] + 1 if chunks else offset
        return dict(**status, items=chunks, next_offset=next_offset,
                    has_more=next_offset < status['chunks'], returned_bytes=used, retrieval_only=False)

    list_chunks = read

    def search(self, attachment_id, owner, query=None, offset=0, limit=10, max_bytes=48_000, *, queries=None, context=0):
        self._page(offset, limit, max_bytes)
        if (query is None) == (queries is None):
            raise ValueError('query和queries需要且只能提供一个')
        terms = [query] if queries is None else queries
        if (not isinstance(terms, list) or not 1 <= len(terms) <= 5
                or any(not isinstance(term, str) or not 1 <= len(term.strip()) <= 500
                       or len(term) > 500 for term in terms)):
            raise ValueError('搜索词需为1至5个非空字面关键词，每个最多500字')
        if type(context) is not int or not 0 <= context <= 2:
            raise ValueError('context需为0至2的相邻块数量')
        terms = list(dict.fromkeys(terms))
        status = self.status(attachment_id, owner)
        patterns = ['%' + term.replace('\\', '\\\\').replace('%', '\\%').replace('_', '\\_') + '%'
                    for term in terms]
        condition = ' OR '.join("text LIKE ? ESCAPE '\\'" for _ in patterns)
        with self.store.connect() as db:
            rows = db.execute('SELECT id,ordinal,text,source FROM document_chunks WHERE attachment_id=? AND ('
                              + condition + ') ORDER BY ordinal LIMIT ? OFFSET ?',
                              (attachment_id, *patterns, limit + 1, offset)).fetchall()
        hits, used = [], 0
        for row in rows[:limit]:
            item = self._chunk(row)
            cost = len(_json(item).encode('utf-8'))
            if used + cost > max_bytes:
                if not hits:
                    raise ValueError('单个命中块超过返回预算，请提高max_bytes后读取')
                break
            hits.append(item)
            used += cost
        hit_ordinals = [item['ordinal'] for item in hits]
        neighbors, context_truncated = [], False
        if context and hits:
            requested = sorted({ordinal + distance for ordinal in hit_ordinals
                                for distance in range(-context, context + 1)
                                if 0 <= ordinal + distance < status['chunks']
                                and ordinal + distance not in hit_ordinals})
            if requested:
                placeholders = ','.join('?' for _ in requested)
                with self.store.connect() as db:
                    nearby = db.execute('SELECT id,ordinal,text,source FROM document_chunks WHERE attachment_id=? AND ordinal IN ('
                                        + placeholders + ') ORDER BY ordinal', (attachment_id, *requested)).fetchall()
                # Spend the remaining budget on the closest neighbors first.
                nearby.sort(key=lambda row: (min(abs(row[1] - hit) for hit in hit_ordinals), row[1]))
                for row in nearby:
                    item = self._chunk(row)
                    cost = len(_json(item).encode('utf-8'))
                    if used + cost > max_bytes:
                        context_truncated = True
                        continue
                    neighbors.append(item)
                    used += cost
                context_truncated = context_truncated or len(nearby) < len(requested)
        neighbors.sort(key=lambda item: item['ordinal'])
        items = sorted(hits + neighbors, key=lambda item: item['ordinal'])
        return dict(**status, items=items, query=query, queries=terms, context=context,
                    hit_ordinals=hit_ordinals, context_ordinals=[item['ordinal'] for item in neighbors],
                    hit_count=len(hits), context_count=len(neighbors), context_truncated=context_truncated,
                    next_offset=offset + len(hits), has_more=len(rows) > len(hits),
                    returned_bytes=used, retrieval_only=True,
                    note='命中和相邻上下文仅用于定位；解析覆盖不代表检索已全文扫描或全量统计')

    def _readable(self, attachment_id, owner):
        record = self._record(attachment_id, owner)
        if record['state'] not in ('ready', 'partial'):
            raise ValueError('请等待文档解析完成；失败或取消状态需显式重试')
        return record

    def iter_chunks(self, attachment_id, owner):
        record = self._readable(attachment_id, owner)
        last_ordinal = -1
        while True:
            # A consumer may persist analysis results between yields. Never hold
            # a SQLite read cursor/transaction while handing a row to it.
            with self.store.connect() as db:
                self._check_generation(record, db)
                rows = db.execute('SELECT id,ordinal,text,source FROM document_chunks WHERE attachment_id=? AND ordinal>? ORDER BY ordinal LIMIT 100',
                                  (attachment_id, last_ordinal)).fetchall()
            if not rows:
                return
            for row in rows:
                last_ordinal = row[1]
                yield {**self._chunk(row), 'coverage_complete': record['state'] == 'ready'}

    @staticmethod
    def _check_generation(record, db):
        current = db.execute('SELECT object_id,parser_version,state,attempt FROM document_parses WHERE attachment_id=? AND owner=?',
                             (record['attachment_id'], record['owner'])).fetchone()
        expected = (record['object_id'], record['parser_version'], record['state'], record['attempt'])
        if current != expected:
            raise ValueError('解析版本或状态已变化，请重新读取文档并创建分析任务')

    def iter_rows(self, attachment_id, owner, table=None):
        record = self._readable(attachment_id, owner)
        if table is not None and (not isinstance(table, str) or len(table) > 500):
            raise ValueError('表名称无效')
        last_name, last_number = '', -1
        while True:
            with self.store.connect() as db:
                self._check_generation(record, db)
                query = ('SELECT table_name,row_number,cells,source FROM document_rows WHERE attachment_id=? '
                         'AND (table_name>? OR (table_name=? AND row_number>?))')
                params = (attachment_id, last_name, last_name, last_number)
                if table is not None:
                    query += ' AND table_name=?'
                    params += (table,)
                query += ' ORDER BY table_name,row_number LIMIT 100'
                rows, size = [], 0
                for row in db.execute(query, params):
                    rows.append(row)
                    size += len(row[2].encode('utf-8')) + len(row[3].encode('utf-8'))
                    if size >= MAX_UNIT_BYTES:
                        break
            if not rows:
                return
            for name, number, cells, source in rows:
                last_name, last_number = name, number
                identifier = hashlib.sha256((attachment_id + ':' + record['object_id'] + ':' + record['parser_version'] + ':' + name + ':' + str(number)).encode()).hexdigest()
                yield dict(id=identifier, table_name=name, row_number=number, cells=json.loads(cells),
                           source=json.loads(source), coverage_complete=record['state'] == 'ready')

    @staticmethod
    def _decimal(value):
        if value is None or str(value).strip() == '':
            return None
        try:
            number = Decimal(str(value).strip())
        except InvalidOperation:
            raise ValueError('数值操作遇到非数值单元格') from None
        if not number.is_finite() or len(number.as_tuple().digits) > 1000 or abs(number.adjusted()) > 1000:
            raise ValueError('数值无效或超过计算范围')
        return number

    def aggregate(self, attachment_id, owner, operation='count', column=None, group_by=None, filters=None, table=None):
        record = self._readable(attachment_id, owner)
        if operation not in ('count', 'sum', 'avg'):
            raise ValueError('数值操作仅支持count、sum、avg')
        if column is not None and (not isinstance(column, str) or not 1 <= len(column) <= 500):
            raise ValueError('列名称无效')
        if operation != 'count' and column is None:
            raise ValueError('sum和avg需要列名称')
        groups = [group_by] if isinstance(group_by, str) else ([] if group_by is None else group_by)
        if not isinstance(groups, list) or len(groups) > 3 or any(not isinstance(key, str) or not 1 <= len(key) <= 500 for key in groups):
            raise ValueError('最多按3个列分组')
        filters = [] if filters is None else filters
        if not isinstance(filters, list) or len(filters) > 10:
            raise ValueError('最多10个筛选条件')
        for item in filters:
            if (not isinstance(item, dict) or set(item) != {'column', 'op', 'value'}
                    or not isinstance(item['column'], str) or not 1 <= len(item['column']) <= 500
                    or item['op'] not in ('eq', 'ne', 'gt', 'gte', 'lt', 'lte', 'contains')
                    or not (item['value'] is None or type(item['value']) in (str, int, float))
                    or len(str(item['value'])) > 1000):
                raise ValueError('筛选条件需column、op、value，支持eq/ne/gt/gte/lt/lte/contains')
        tables = record['metadata'].get('tables', [])
        selected = [item for item in tables if table is None or item['name'] == table]
        if not selected:
            raise ValueError('文档没有对应的结构化表格')
        required = groups + ([column] if column is not None else []) + [item['column'] for item in filters]
        for item in selected:
            if any(key not in item['columns'] for key in required):
                raise ValueError('指定列不在所选表格中，请指定table并核对列名称')
        results, examined, matched, invalid, examples = {}, 0, 0, 0, []
        # Decimal sum is exact within the bounded source precision. avg explicitly
        # uses 34 significant decimal digits, since many averages do not terminate.
        with localcontext() as arithmetic:
            # Accepted values span adjusted exponents -1000..1000 with up to
            # 1000 digits. Their sum may therefore need ~3007 digits, including
            # carry from the maximum 500,000 input units.
            arithmetic.prec = 3200
            for row in self.iter_rows(attachment_id, owner, table):
                examined += 1
                cells = row['cells']
                accepted = True
                for item in filters:
                    current, value, op = cells.get(item['column']), item['value'], item['op']
                    if op in ('gt', 'gte', 'lt', 'lte'):
                        left, right = self._decimal(current), self._decimal(value)
                        if left is None or right is None:
                            accepted = False
                        else:
                            accepted = {'gt': left > right, 'gte': left >= right, 'lt': left < right, 'lte': left <= right}[op]
                    elif op == 'contains':
                        accepted = str(value) in str(current or '')
                    elif op == 'eq':
                        accepted = current == (None if value is None else str(value))
                    else:
                        accepted = current != (None if value is None else str(value))
                    if not accepted:
                        break
                if not accepted:
                    continue
                matched += 1
                key = tuple(cells.get(name) for name in groups)
                if key not in results:
                    if len(results) >= 1000:
                        raise ValueError('分组超过1000项，请缩小筛选范围')
                    results[key] = dict(count=0, sum=Decimal(0), empty=0, invalid=0)
                bucket = results[key]
                value = cells.get(column) if column is not None else '1'
                if value is None or str(value).strip() == '':
                    bucket['empty'] += 1
                    continue
                if operation == 'count':
                    bucket['count'] += 1
                    continue
                try:
                    numeric = self._decimal(value)
                except ValueError:
                    invalid += 1
                    bucket['invalid'] += 1
                    if len(examples) < 5:
                        examples.append(dict(value=str(value)[:200], source=row['source']))
                    continue
                bucket['count'] += 1
                bucket['sum'] += numeric
            if not results and not groups:
                results[()] = dict(count=0, sum=Decimal(0), empty=0, invalid=0)
            output = []
            for key, bucket in results.items():
                if operation == 'count':
                    value = str(bucket['count'])
                elif operation == 'sum':
                    value = format(bucket['sum'], 'f')
                elif bucket['count']:
                    with localcontext() as average_context:
                        average_context.prec = 34
                        value = str(bucket['sum'] / bucket['count'])
                else:
                    value = None
                output.append(dict(group=dict(zip(groups, key)), value=value, rows=bucket['count'],
                                   empty_rows=bucket['empty'], non_numeric_rows=bucket['invalid']))
        coverage = self._public(record)['coverage']
        return dict(attachment_id=attachment_id, object_id=record['object_id'], parser_version=record['parser_version'],
                    operation=operation, column=column, table=table, group_by=groups, filters=filters,
                    groups=output, examined_rows=examined, matched_rows=matched, non_numeric_rows=invalid,
                    non_numeric_examples=examples, complete=coverage['complete'] and invalid == 0, coverage=coverage,
                    numeric_policy='Decimal; sum exact, avg 34 significant digits; empty and nonnumeric values explicitly counted',
                    warnings=self._public(record)['warnings'])
